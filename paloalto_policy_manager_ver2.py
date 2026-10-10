"""
paloalto_policy_manager.py - Reads a list of requested firewall policies
from Excel and runs them through this pipeline for each row:

    Excel -> Resolve zones from routing -> Validate objects -> Check
    existing rules -> Check duplicate/merge/shadowing -> Show proposed
    change -> DRY RUN -> Apply

INVENTORY AND CREDENTIALS COME FROM ANSIBLE AWX
--------------------------------------------------
This script no longer stores a firewall hostname, API key, username or
password anywhere - not in the script, not in paloalto_config.json.
AWX owns both:

    - INVENTORY: the firewall(s) live in an AWX inventory. The AWX
      playbook (awx_run_policy_manager.yml) runs this script once per
      firewall host and passes that host's address in the PANOS_HOSTNAME
      environment variable.
    - CREDENTIALS: a custom AWX credential type ("Palo Alto PAN-OS API")
      injects the secret into the job as environment variables:
          PANOS_API_KEY                     (preferred), or
          PANOS_USERNAME + PANOS_PASSWORD   (SDK generates a key)

The script only READS those environment variables (see
load_awx_connection()). If any of hostname/api_key/username/password is
still sitting in paloalto_config.json from an older version, it is
IGNORED and a warning is printed - delete it from the file.

For a quick local test outside AWX, just set the same environment
variables yourself before running, e.g. (PowerShell):
    $env:PANOS_HOSTNAME="10.0.0.1"; $env:PANOS_API_KEY="..."

You do NOT provide Source Zone / Destination Zone in the Excel sheet.
Instead, for every Source and Destination address you give it, the
script asks the firewall's own routing table (via a FIB lookup) which
interface that traffic would use, then looks up which zone that
interface belongs to - the same two steps ("test routing fib-lookup",
then "show interface ... | match Zone") a network engineer would do by
hand. See STEP 0 below for exactly how this works and its limits.

WHAT THIS SCRIPT DOES NOT DO
-----------------------------
It does not look at traffic logs or try to figure out what's currently
being blocked in practice. Its "shadowing" check (see below) is a
config-level check against the existing rulebase, not a traffic-log
analysis - a deliberate scope limit to keep this predictable.

Zone resolution (STEP 0) is a routing-table lookup, not a guarantee -
see the limits called out in that section. If your network isn't
symmetric (the return path differs from the forward path), the zone
this script infers for an address may not match reality - review the
proposed change before applying.

ADDRESS NORMALIZATION (fixes duplicate rules for equivalent subnets)
------------------------------------------------------------------------
Any Source/Destination value written as an IP/CIDR is normalized to its
canonical network form before it's ever compared or stored - e.g.
"50.50.50.50/24" and "50.50.50.0/24" are the exact same /24 network, so
both get normalized to "50.50.50.0/24". "any", address ranges, and
existing Address object/group names are left exactly as typed. See
normalize_address().

Address merging is subnet-aware (see merge_addresses()): if a rule
already has a broader network and a narrower request comes in that's
already inside it, nothing is added; if a broader network comes in that
now covers an existing narrower entry, that narrower entry is replaced.
This applies to BOTH Source and Destination merges.

DEFAULT-ROUTE EXCLUSION (never infers a zone from the default route)
------------------------------------------------------------------------
A FIB lookup returns whichever route is the BEST match for an IP - and
if nothing more specific exists, that's the default route (0.0.0.0/0),
which says nothing about which zone an address belongs to. Before
trusting a FIB-lookup match, STEP 0 cross-checks the same virtual
router's routing table ("show routing route") and confirms the most
specific route covering that IP has a prefix length greater than 0. If
the only cover is the default route, that virtual router is treated as
having no route. If every VR ends up that way, the row is marked "Could
not implement - no specific route found" and NO rule is created. See
get_virtual_router_routes() and most_specific_route_prefixlen().

TAG AUTO-CREATION
-------------------
PAN-OS rejects a rule whose Tag value doesn't reference an existing Tag
object. If your Tag column has a value that doesn't exist yet, this
script creates it (only when apply is on). Tags are ONLY set when
creating a brand-new rule - merging into an existing rule never touches
that rule's tag.

SETUP
-----
1. In the AWX Execution Environment (or locally):
       pip install pan-os-python openpyxl

2. paloalto_config.json (OPTIONAL now - next to this file) only holds
   run-mode settings. No connection details:

    {
      "excel_path": "requests.xlsx",
      "request_sheet": "Request",
      "implementation_sheet": "Implementation",
      "apply": false,
      "commit": false
    }

   Add "virtual_router": "<name>" only if you want to FORCE a single
   virtual router (see STEP 0). If the file doesn't exist at all, the
   defaults above are used.

   From a terminal or the AWX playbook, --apply / --commit / --excel /
   --request-sheet / --implementation-sheet override the config file.

PACKAGING AS A .EXE
--------------------
Still works the same (pyinstaller --onefile paloalto_policy_manager.py);
get_base_dir() makes config/Excel lookups relative to the .exe's real
folder. The .exe also reads its connection from the PANOS_* environment
variables - it has no other way to get credentials now.

3. Your Excel file needs a "Request" sheet with a header row containing
   these columns (order doesn't matter - read by column name):

      Tag | Source | Destination | Destination Port | Application |
      Service | Action

   NOTE on Destination Port vs Service: if you fill in Destination Port,
   the PAN-OS service is built from Destination Port + Service together -
   Service, if given, just supplies the protocol ("tcp"/"udp"; default
   tcp). E.g. Destination Port "443" becomes tcp/443. If Destination Port
   is blank, Service works as a bare port, "tcp/443", "any",
   "application-default", or an existing Service object name. See
   build_raw_service_values().

   NOTE on "Application": blank/missing defaults to "any".
   "application-default" typed here is replaced by "any" (it's only
   valid in the Service field).

   Source / Destination / Destination Port / Service / Application can
   each hold more than one comma-separated value.

STEP 0 - RESOLVE ZONES FROM ROUTING
--------------------------------------
For every Source and Destination address in a row, the script:
    1. Turns the address into a single IP to look up (raw IP/CIDR, or an
       Address object's underlying IP; "any" is skipped).
    2. Runs a FIB lookup against EVERY virtual router (unless
       "virtual_router" is forced in the config). If VRs disagree on the
       zone, that's flagged rather than guessed.
    3. Looks up which zone that interface belongs to.
    A default-route-only match never counts (see above).

If a row's Sources (or Destinations) resolve to DIFFERENT zones, that's
flagged ("split into separate rows") - one rule has one set of zones.

LIMITS: Address GROUPS can't be auto-resolved; asymmetric routing can
give the wrong zone; many IPs x many VRs = many API calls (cached per run).

STEP 1 - VALIDATE OBJECTS
---------------------------
Each row is checked for: zones resolved; addresses look like IP/CIDR/
range, a known Address object/group, or "any"; Destination Port is
numeric; non-port Service values exist on the firewall; Action is a real
PAN-OS action. A failing row is marked "Validation error" and skipped.

STEP 2 - CHECK EXISTING RULES / STEP 3 - DUPLICATE, MERGE OR SHADOWING
------------------------------------------------------------------------
An existing rule is a CANDIDATE when these match exactly:
    - Source Zone and Destination Zone (as resolved in STEP 0)
    - Application
    - Action

Then the three "traffic" fields are compared - Source, Destination and
Service (Service is where Destination Port ends up, e.g. svc-tcp-8443):

    DUPLICATE - a candidate rule already COVERS all three (exact match,
        "any", or a broader subnet containing the request). Nothing to do.

    MERGE - a candidate rule has EXACTLY the same values on two of the
        three fields, and only the third differs. The differing value is
        merged into that rule instead of creating a new one:
          * Same Destination + same Port, new Source
                -> Source added to the existing rule
          * Same Source + same Port, new Destination
                -> Destination added to the existing rule
          * Same Source + same Destination, new Port/Service
                -> Service (port) added to the existing rule
        Checked in that order (Source first) if more than one would fit.

    Why only ONE field may differ: a rule allows every combination of
    its sources x destinations x services. If only one field differs,
    merging adds exactly the requested traffic and nothing else. If two
    fields differ, merging would also open combinations nobody asked
    for (e.g. the old source to the new destination), so a NEW rule is
    created instead. Likewise the two "same" fields must be EXACTLY
    equal - a rule for 10.1.1.0/24 is not merged into for a request from
    10.1.1.5 to a new destination, because that would open the new
    destination to the whole /24.

    Service merges are skipped (new rule instead) when either side uses
    "application-default", since PAN-OS doesn't allow mixing that
    keyword with port-based services on one rule.

If no candidate fits, the script checks for SHADOWING: does an earlier
rule cover this exact traffic with a DIFFERENT action? If so, the row is
marked "Shadow conflict - not created" and nothing is created.

Everything that decides this is in find_matching_rule(),
rule_base_matches() and check_shadowing().

STEP 4 - SHOW PROPOSED CHANGE
-------------------------------
Whatever the outcome, the script prints what it found / would do for
that row BEFORE doing anything, and writes the same into the
Implementation sheet's Outcome/Details columns.

STEP 5 - DRY RUN / APPLY
---------------------------
Default is a DRY RUN - nothing changes on the firewall. "apply": true
(or --apply) stages changes in the candidate config; "commit": true (or
--commit) also commits.

BATCH SIMULATION IN DRY RUN: rows are compared against everything
proposed so far in the same run, in sheet order - so a later row that
would duplicate or merge into a rule proposed by an earlier row is
reported correctly even in a dry run.
"""

import os
import re
import sys
import json
import argparse
import ipaddress
from pathlib import Path
from datetime import datetime

import openpyxl
from panos.firewall import Firewall
from panos.policies import Rulebase, SecurityRule
from panos.objects import ServiceObject, AddressObject, AddressGroup, Tag
from panos.network import Zone, VirtualRouter


def get_base_dir() -> Path:
    """The folder this script (or the packaged .exe) actually lives in -
    used to find paloalto_config.json and, by default, requests.xlsx.
    Under PyInstaller, __file__ points at a temp extraction folder, so
    sys.executable's folder is used instead when sys.frozen is set."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).parent


CONFIG_FILE = get_base_dir() / "paloalto_config.json"

REQUEST_COLUMNS = [
    "Tag", "Source", "Destination", "Destination Port", "Application", "Service", "Action",
]

VALID_ACTIONS = {"allow", "deny", "drop", "reset-client", "reset-server", "reset-both"}

# Loose IP / CIDR / range matcher - tells "raw address" apart from
# "object name". A sanity check, not a full validator.
IP_LIKE_PATTERN = re.compile(r"^\d{1,3}(\.\d{1,3}){0,3}(/\d{1,2})?(-\d{1,3}(\.\d{1,3}){0,3})?$")

# The three "traffic" fields the duplicate/merge logic compares.
# Destination Port lives inside "service" (e.g. svc-tcp-8443).
MERGEABLE_FIELDS = ("source", "destination", "service")


def normalize_address(value: str) -> str:
    """Normalizes an IP/CIDR to its canonical network form, so e.g.
    '50.50.50.50/24' and '50.50.50.0/24' compare equal. "any", ranges and
    object/group names are returned unchanged."""
    text = value.strip()
    if text.lower() == "any":
        return "any"
    if "/" in text and "-" not in text:
        try:
            return str(ipaddress.ip_network(text, strict=False))
        except ValueError:
            return text  # not a valid CIDR - validate_row() will flag it
    return text


def normalize_address_list(values: list) -> list:
    return [normalize_address(v) for v in values]


# --- Config / connection -------------------------------------------------

def load_config() -> dict:
    """Loads run-mode settings from paloalto_config.json (optional).
    Connection details are NOT read from here any more - see
    load_awx_connection()."""
    config = {}
    if CONFIG_FILE.exists():
        config = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))

    for leftover in ("hostname", "api_key", "username", "password"):
        if config.get(leftover):
            print(f"WARNING: '{leftover}' in {CONFIG_FILE.name} is ignored - the firewall and credentials "
                  f"now come from AWX (PANOS_* environment variables). Please remove it from the file.")

    config.setdefault("excel_path", "requests.xlsx")
    config.setdefault("request_sheet", "Request")
    config.setdefault("implementation_sheet", "Implementation")
    config.setdefault("apply", False)
    config.setdefault("commit", False)
    config.setdefault("virtual_router", None)  # None = check every virtual router
    return config


def load_awx_connection() -> dict:
    """Reads the firewall address (from the AWX inventory host) and the
    credential (from the AWX custom credential type) out of environment
    variables. Nothing secret is ever stored in this script or its
    config file."""
    hostname = os.environ.get("PANOS_HOSTNAME", "").strip()
    api_key = os.environ.get("PANOS_API_KEY", "").strip()
    username = os.environ.get("PANOS_USERNAME", "").strip()
    password = os.environ.get("PANOS_PASSWORD", "")  # not stripped - spaces can be part of a password

    if not hostname:
        print("ERROR: PANOS_HOSTNAME is not set. In AWX this comes from the inventory host "
              "(see awx_run_policy_manager.yml). For a local test, set it yourself.")
        sys.exit(1)
    if not api_key and not (username and password):
        print("ERROR: no credential found. Attach the 'Palo Alto PAN-OS API' credential to the AWX job "
              "template, so it injects PANOS_API_KEY (or PANOS_USERNAME + PANOS_PASSWORD).")
        sys.exit(1)

    return {"hostname": hostname, "api_key": api_key, "username": username, "password": password}


def connect_to_firewall(connection: dict, config: dict):
    """Connects to the firewall and pulls back everything the validation,
    zone-resolution, and matching steps need. Returns (firewall_object,
    rulebase_object, list_of_existing_rules, dict_of_existing_services_by_name,
    set_of_known_tag_names, set_of_zone_names, set_of_known_address_names,
    dict_of_address_object_values, list_of_virtual_router_names_to_check)."""
    if connection["api_key"]:
        fw = Firewall(connection["hostname"], api_key=connection["api_key"])
    else:
        fw = Firewall(connection["hostname"], connection["username"], connection["password"])

    rulebase = fw.add(Rulebase())
    existing_rules = SecurityRule.refreshall(rulebase)

    existing_services = ServiceObject.refreshall(fw)
    existing_services_by_name = {s.name: s for s in existing_services}

    known_tag_names = {t.name for t in Tag.refreshall(fw)}

    zone_names = {z.name for z in Zone.refreshall(fw)}

    address_objects = AddressObject.refreshall(fw)
    address_groups = AddressGroup.refreshall(fw)
    known_address_names = {a.name for a in address_objects} | {g.name for g in address_groups}
    # Address groups deliberately excluded - see STEP 0 limits.
    address_object_values = {a.name: a.value for a in address_objects}

    all_virtual_router_names = [vr.name for vr in VirtualRouter.refreshall(fw)]
    if not all_virtual_router_names:
        print("ERROR: no virtual routers found on this firewall - can't do routing-based zone lookups.")
        sys.exit(1)

    forced_virtual_router = config.get("virtual_router")
    if forced_virtual_router:
        if forced_virtual_router not in all_virtual_router_names:
            print(f"ERROR: virtual_router '{forced_virtual_router}' not found on the firewall. Available: {all_virtual_router_names}")
            sys.exit(1)
        virtual_routers_to_check = [forced_virtual_router]
    else:
        virtual_routers_to_check = all_virtual_router_names

    return fw, rulebase, existing_rules, existing_services_by_name, known_tag_names, zone_names, known_address_names, address_object_values, virtual_routers_to_check


# --- STEP 0: Resolving zones from the routing table -----------------------

def get_ip_for_lookup(value: str, address_object_values: dict):
    """Turns one Source/Destination value into a single IP for a FIB
    lookup. Returns (ip_or_None, error_or_None); ("any", None) for 'any'."""
    text = value.strip()
    if text.lower() == "any":
        return ("any", None)

    if IP_LIKE_PATTERN.match(text):
        base = text.split("/")[0].split("-")[0]
        return (base, None)

    if text in address_object_values:
        object_value = address_object_values.get(text)
        if not object_value:
            return (None, f"Address object '{text}' has no value set")
        base = str(object_value).split("/")[0].split("-")[0]
        return (base, None)

    return (None, f"'{text}' is an Address group or unknown object - can't auto-resolve its zone, specify manually")


def fib_lookup_interface(fw, ip: str, virtual_router: str):
    """Which interface would the firewall route this IP through? None if no route."""
    cmd = f"<test><routing><fib-lookup><virtual-router>{virtual_router}</virtual-router><ip>{ip}</ip></fib-lookup></routing></test>"
    try:
        response = fw.op(cmd, cmd_xml=False)
    except Exception as e:
        raise RuntimeError(f"FIB lookup for {ip} failed: {e}")
    interface_elem = response.find("./result/interface")
    if interface_elem is None or not interface_elem.text:
        return None
    return interface_elem.text.strip()


def get_zone_for_interface(fw, interface_name: str, interface_zone_cache: dict):
    """Which zone is this interface in? Cached per run."""
    if interface_name in interface_zone_cache:
        return interface_zone_cache[interface_name]
    cmd = f"<show><interface>{interface_name}</interface></show>"
    try:
        response = fw.op(cmd, cmd_xml=False)
    except Exception as e:
        raise RuntimeError(f"'show interface {interface_name}' failed: {e}")
    zone_elem = response.find("./result/ifnet/zone")
    zone = zone_elem.text.strip() if zone_elem is not None and zone_elem.text else None
    interface_zone_cache[interface_name] = zone
    return zone


def get_virtual_router_routes(fw, virtual_router: str, route_cache: dict):
    """All routes in one VR's table ("show routing route"), cached per run.
    Used only to tell a specific-route match from a default-route match."""
    if virtual_router in route_cache:
        return route_cache[virtual_router]
    cmd = f"<show><routing><route><virtual-router>{virtual_router}</virtual-router></route></routing></show>"
    try:
        response = fw.op(cmd, cmd_xml=False)
    except Exception as e:
        raise RuntimeError(f"'show routing route' for {virtual_router} failed: {e}")
    routes = []
    for entry in response.findall("./result/entry"):
        dest_elem = entry.find("destination")
        if dest_elem is None or not dest_elem.text:
            continue
        try:
            routes.append(ipaddress.ip_network(dest_elem.text.strip(), strict=False))
        except ValueError:
            continue
    route_cache[virtual_router] = routes
    return routes


def most_specific_route_prefixlen(routes: list, ip: str):
    """Prefix length of the most specific route containing this IP, or
    None if nothing covers it. 0 means default-route only."""
    try:
        address = ipaddress.ip_address(ip)
    except ValueError:
        return None
    matching = [net for net in routes if address in net]
    if not matching:
        return None
    return max(net.prefixlen for net in matching)


def resolve_zone_for_value(fw, value: str, virtual_routers_to_check: list, address_object_values: dict, interface_zone_cache: dict, ip_zone_cache: dict, route_cache: dict):
    """Resolves ONE Source/Destination value to a single zone, checking
    every VR in virtual_routers_to_check. Returns (zone_or_None, error_or_None)."""
    ip, error = get_ip_for_lookup(value, address_object_values)
    if error:
        return (None, error)
    if ip == "any":
        return ("any", None)

    if ip in ip_zone_cache:
        return ip_zone_cache[ip]

    zone_by_vr = {}
    default_route_only_vrs = []
    for vr in virtual_routers_to_check:
        try:
            routes = get_virtual_router_routes(fw, vr, route_cache)
        except RuntimeError:
            routes = None  # couldn't read the table - fall back to trusting the FIB lookup
        if routes is not None:
            prefixlen = most_specific_route_prefixlen(routes, ip)
            if prefixlen == 0:
                default_route_only_vrs.append(vr)
                continue
        try:
            interface_name = fib_lookup_interface(fw, ip, vr)
        except RuntimeError:
            continue
        if interface_name is None:
            continue
        try:
            zone = get_zone_for_interface(fw, interface_name, interface_zone_cache)
        except RuntimeError:
            continue
        if zone:
            zone_by_vr[vr] = zone

    if not zone_by_vr:
        if default_route_only_vrs:
            result = (None, f"{value} ({ip}) only matches the DEFAULT route in {', '.join(default_route_only_vrs)} - not specific enough to confidently resolve a zone; add a specific route or specify the zone manually")
        else:
            result = (None, f"no route found for {value} ({ip}) in any virtual router checked ({', '.join(virtual_routers_to_check)})")
        ip_zone_cache[ip] = result
        return result

    distinct_zones = set(zone_by_vr.values())
    if len(distinct_zones) > 1:
        breakdown = ", ".join(f"{vr}->{zone}" for vr, zone in zone_by_vr.items())
        result = (None, f"{value} ({ip}) routes to different zones depending on virtual router ({breakdown}) - set \"virtual_router\" in paloalto_config.json to pick one")
        ip_zone_cache[ip] = result
        return result

    result = (distinct_zones.pop(), None)
    ip_zone_cache[ip] = result
    return result


def resolve_zone_for_list(fw, values: list, virtual_routers_to_check: list, address_object_values: dict, interface_zone_cache: dict, ip_zone_cache: dict, route_cache: dict):
    """Resolves a whole Source or Destination cell to one zone. Returns
    (zone_or_None, list_of_problem_strings)."""
    zones_found = set()
    problems = []

    for value in values:
        zone, error = resolve_zone_for_value(fw, value, virtual_routers_to_check, address_object_values, interface_zone_cache, ip_zone_cache, route_cache)
        if error:
            problems.append(error)
        else:
            zones_found.add(zone)

    if problems:
        return (None, problems)
    if len(zones_found) > 1:
        return (None, [f"resolves to multiple different zones {sorted(zones_found)} - split into separate rows"])
    if not zones_found:
        return (None, ["no addresses to resolve a zone from"])
    return (zones_found.pop(), [])


# --- Resolving Service cell values into valid PAN-OS service names --------
#
# Accepted formats in the Service column:
#   any / application-default -> used as-is
#   443 or tcp/443            -> object "svc-tcp-443" (auto-created)
#   udp/53                    -> object "svc-udp-53"  (auto-created)
#   service-https             -> assumed to be an existing object name

PORT_PATTERN = re.compile(r"^(tcp|udp)?[/\s]?(\d+)$", re.IGNORECASE)


def parse_service_entry(raw_value: str):
    """Returns ("keyword", kw) / ("port", protocol, port) / ("name", raw_value)."""
    text = raw_value.strip()
    if text.lower() in ("any", "application-default"):
        return ("keyword", text.lower())

    match = PORT_PATTERN.match(text)
    if match:
        protocol = (match.group(1) or "tcp").lower()
        port = match.group(2)
        return ("port", protocol, port)

    return ("name", text)


def resolve_service_names(raw_service_values: list, existing_services_by_name: dict, taken_service_names: set, fw, apply_changes: bool):
    """Turns raw Service values into names valid on a security rule,
    creating Service objects for bare ports. Returns (resolved_names, notes)."""
    resolved_names = []
    notes = []

    for raw_value in raw_service_values:
        kind, *rest = parse_service_entry(raw_value)

        if kind in ("keyword", "name"):
            resolved_names.append(rest[0])
            continue

        protocol, port = rest
        object_name = f"svc-{protocol}-{port}"
        resolved_names.append(object_name)

        if object_name in existing_services_by_name or object_name in taken_service_names:
            continue

        taken_service_names.add(object_name)
        notes.append(f"service object '{object_name}' ({protocol}/{port})")
        if apply_changes:
            new_service = ServiceObject(name=object_name, protocol=protocol, destination_port=port)
            fw.add(new_service)
            new_service.create()

    return resolved_names, notes


# --- Resolving Tag cell values into valid PAN-OS tag references ------------

def resolve_tags(raw_tag_values: list, known_tag_names: set, taken_tag_names: set, fw, apply_changes: bool):
    """Auto-creates any Tag object that doesn't exist yet. Returns
    (tag_names, notes)."""
    notes = []
    for tag_name in raw_tag_values:
        if tag_name in known_tag_names or tag_name in taken_tag_names:
            continue
        taken_tag_names.add(tag_name)
        notes.append(f"tag '{tag_name}'")
        if apply_changes:
            new_tag = Tag(name=tag_name)
            fw.add(new_tag)
            new_tag.create()
    return raw_tag_values, notes


def build_raw_service_values(service_cell, dest_port_cell) -> list:
    """Combines Service and Destination Port into raw service values.
    If Destination Port is set, it's authoritative and Service supplies
    the protocol by position (default tcp); extra Service entries are
    kept as-is. If Destination Port is empty, Service is used as before."""
    ports = split_values(dest_port_cell)
    services = split_values(service_cell)

    if not ports:
        return services

    combined = []
    for i, port in enumerate(ports):
        protocol = "tcp"
        if i < len(services) and services[i].lower() in ("tcp", "udp"):
            protocol = services[i].lower()
        combined.append(f"{protocol}/{port}")

    if len(services) > len(ports):
        combined.extend(services[len(ports):])

    return combined


# --- Reading the request sheet -------------------------------------------

def split_values(cell_value) -> list:
    """'10.1.1.5, 10.1.1.6' -> ['10.1.1.5', '10.1.1.6']"""
    if cell_value is None:
        return []
    text = str(cell_value).strip()
    if not text:
        return []
    return [part.strip() for part in text.split(",") if part.strip()]


def read_requests(excel_path: Path, sheet_name: str) -> list:
    """Reads the request sheet into a list of dicts keyed by header."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: sheet '{sheet_name}' not found in {excel_path}. Sheets present: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[sheet_name]

    header_row = [cell.value for cell in ws[1]]
    rows = []
    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(header_row, row_cells))
        if all(v is None or str(v).strip() == "" for v in row_dict.values()):
            continue
        rows.append(row_dict)
    return rows


# --- STEP 1: Validating objects -------------------------------------------

def looks_like_address(value: str, known_address_names: set) -> bool:
    if value.lower() == "any":
        return True
    if IP_LIKE_PATTERN.match(value):
        return True
    return value in known_address_names


def validate_row(req: dict, raw_service_values: list, known_address_names: set, existing_services_by_name: dict) -> list:
    """Returns a list of problems with one request row (empty = OK)."""
    problems = []

    source_addrs = split_values(req.get("Source"))
    dest_addrs = split_values(req.get("Destination"))
    if not source_addrs:
        problems.append("Source is empty")
    if not dest_addrs:
        problems.append("Destination is empty")
    for addr in source_addrs + dest_addrs:
        if not looks_like_address(addr, known_address_names):
            problems.append(f"'{addr}' doesn't look like a valid IP/CIDR and isn't a known Address object")

    dest_ports = split_values(req.get("Destination Port"))
    for port in dest_ports:
        if not port.isdigit():
            problems.append(f"Destination Port '{port}' isn't a number")

    if not raw_service_values:
        problems.append("Service (or Destination Port) is empty")
    for value in raw_service_values:
        kind, *_ = parse_service_entry(value)
        if kind == "name" and value not in existing_services_by_name:
            problems.append(f"service '{value}' isn't a bare port and isn't a known Service object on the firewall")

    action = str(req.get("Action") or "").strip().lower()
    if not action:
        problems.append("Action is empty")
    elif action not in VALID_ACTIONS:
        problems.append(f"Action '{action}' isn't a valid PAN-OS action ({', '.join(sorted(VALID_ACTIONS))})")

    return problems


# --- Coverage helpers -------------------------------------------------------

def covers(rule_field_values, req_values: set) -> bool:
    """True if the rule's field covers every requested value ('any'/empty
    covers everything). Exact string match - for zones/service/app."""
    rule_set = set(rule_field_values or [])
    if not rule_set or "any" in rule_set:
        return True
    return req_values.issubset(rule_set)


def covers_addresses(rule_field_values, req_values: set) -> bool:
    """Like covers(), but understands subnet containment (a rule with
    10.1.1.0/24 covers a request for 10.1.1.5)."""
    rule_values = normalize_address_list(rule_field_values or [])
    rule_set = set(rule_values)
    if not rule_set or "any" in rule_set:
        return True
    if req_values.issubset(rule_set):
        return True
    for req_value in req_values:
        contained = False
        req_network = _as_network_or_none(req_value)
        for rule_value in rule_set:
            if req_value == rule_value:
                contained = True
                break
            if req_network is not None:
                rule_network = _as_network_or_none(rule_value)
                if rule_network is None:
                    continue
                try:
                    if req_network.subnet_of(rule_network):
                        contained = True
                        break
                except TypeError:
                    continue  # IPv4 vs IPv6 - not comparable
        if not contained:
            return False
    return True


def _as_network_or_none(value: str):
    try:
        return ipaddress.ip_network(value, strict=False)
    except ValueError:
        return None


# --- Merging values into an existing rule -------------------------------------

def merge_addresses(current_values: list, new_values: list):
    """Merges new Source OR Destination address(es) into an existing
    rule's list, subnet-aware: a value already covered by a broader
    network isn't added; a new broader network replaces narrower entries
    it now covers. Returns (merged_list, change_descriptions) - empty
    descriptions means nothing changed."""
    merged = list(current_values)
    changes = []

    for new_value in new_values:
        if new_value in merged:
            continue

        new_network = _as_network_or_none(new_value)

        if new_network is not None:
            already_covered_by = None
            for existing in merged:
                existing_network = _as_network_or_none(existing)
                if existing_network is None:
                    continue
                try:
                    if new_network.subnet_of(existing_network):
                        already_covered_by = existing
                        break
                except TypeError:
                    continue
            if already_covered_by is not None:
                changes.append(f"'{new_value}' already covered by existing '{already_covered_by}' - not added")
                continue

            redundant = []
            for existing in merged:
                existing_network = _as_network_or_none(existing)
                if existing_network is None:
                    continue
                try:
                    if existing_network.subnet_of(new_network) and existing != new_value:
                        redundant.append(existing)
                except TypeError:
                    continue
            for r in redundant:
                merged.remove(r)
                changes.append(f"removed '{r}' - now covered by broader '{new_value}'")

        merged.append(new_value)
        changes.append(f"added '{new_value}'")

    # A list of only "already covered" notes means nothing actually changed.
    if not any(c.startswith(("added", "removed")) for c in changes):
        return merged, []
    return merged, changes


def merge_service_names(current_services: list, new_services: list):
    """Merges new service name(s) (ports become svc-<proto>-<port>
    objects before this point) into an existing rule's service list.
    Returns (merged_list, change_descriptions)."""
    if "any" in new_services:
        if list(current_services) == ["any"]:
            return list(current_services), []
        return ["any"], [f"replaced {current_services} with 'any'"]

    merged = list(current_services)
    changes = []
    for service in new_services:
        if service not in merged:
            merged.append(service)
            changes.append(f"added '{service}'")
    return merged, changes


# --- STEP 2/3: Matching an existing rule (duplicate or merge) -----------------

def _rule_field(rule, field: str) -> list:
    """A rule's Source/Destination (normalized) or Service values."""
    if field == "source":
        return normalize_address_list(rule.source or [])
    if field == "destination":
        return normalize_address_list(rule.destination or [])
    return list(rule.service or [])


def _field_covers(field: str, rule_values: list, req_values: set) -> bool:
    if field in ("source", "destination"):
        return covers_addresses(rule_values, req_values)
    return covers(rule_values, req_values)


def rule_base_matches(rule, req_action: str, req_application: set, req_source_zone: str, req_dest_zone: str) -> bool:
    """Zones, Application and Action must match exactly before a rule is
    even considered for duplicate/merge."""
    return (
        set(rule.fromzone or []) == {req_source_zone}
        and set(rule.tozone or []) == {req_dest_zone}
        and set(rule.application or []) == req_application
        and (rule.action or "").strip().lower() == req_action
    )


def find_matching_rule(existing_rules: list, req: dict, resolved_service_names: list, req_application_list: list, req_source_zone: str, req_dest_zone: str, req_source_set: set, req_dest_set: set):
    """Returns (rule, merge_field):
        (rule, None)          - DUPLICATE: rule already covers the request
        (rule, "source")      - same Destination + Service, merge new Source
        (rule, "destination") - same Source + Service, merge new Destination
        (rule, "service")     - same Source + Destination, merge new Port/Service
        (None, None)          - nothing fits, go to shadow check / new rule
    See STEP 2/3 in the module docstring for why only ONE field may differ."""
    req_action = str(req.get("Action") or "").strip().lower()
    req_application = set(req_application_list)
    requested = {
        "source": req_source_set,
        "destination": req_dest_set,
        "service": set(resolved_service_names),
    }

    candidates = [r for r in existing_rules if rule_base_matches(r, req_action, req_application, req_source_zone, req_dest_zone)]

    # Pass 1: is this traffic already fully allowed by some candidate?
    for rule in candidates:
        if all(_field_covers(f, _rule_field(rule, f), requested[f]) for f in MERGEABLE_FIELDS):
            return rule, None

    # Pass 2: exactly one field differs, the other two are EXACTLY equal.
    for field in MERGEABLE_FIELDS:
        others = [f for f in MERGEABLE_FIELDS if f != field]
        for rule in candidates:
            if not all(set(_rule_field(rule, f)) == requested[f] for f in others):
                continue
            if field == "service":
                combined = set(_rule_field(rule, "service")) | requested["service"]
                if "application-default" in combined:
                    continue  # PAN-OS can't mix application-default with ports
            return rule, field

    return None, None


# --- STEP 3 (no-match branch): checking for shadowing ----------------------

def check_shadowing(existing_rules: list, req: dict, resolved_service_names: list, req_application_list: list, req_source_zone: str, req_dest_zone: str, req_source_addr: set, req_dest_addr: set):
    """Finds an existing rule that matches all of this request's traffic
    but with a DIFFERENT action. Returns that rule or None."""
    req_service = set(resolved_service_names)
    req_application = set(req_application_list)
    req_action = str(req.get("Action") or "").strip().lower()

    for rule in existing_rules:
        rule_action = (rule.action or "").strip().lower()
        if rule_action == req_action:
            continue
        if (
            covers(rule.fromzone, {req_source_zone})
            and covers_addresses(rule.source, req_source_addr)
            and covers(rule.tozone, {req_dest_zone})
            and covers_addresses(rule.destination, req_dest_addr)
            and covers(rule.service, req_service)
            and covers(rule.application, req_application)
        ):
            return rule
    return None


# --- Building the new-rule name -------------------------------------------

def make_rule_name(req: dict, dest_zone: str, taken_names: set) -> str:
    """Tag + Destination Zone + Action, de-duplicated with a counter."""
    tag = str(req.get("Tag") or "rule").strip().replace(" ", "-")
    dest_zone = str(dest_zone or "any").strip().replace(" ", "-")
    action = str(req.get("Action") or "action").strip().replace(" ", "-")
    base_name = f"{tag}-{dest_zone}-{action}"

    name = base_name
    counter = 2
    while name in taken_names:
        name = f"{base_name}-{counter}"
        counter += 1
    return name


# --- Processing one request row -------------------------------------------

def normalize_application_list(raw_values: list):
    """Replaces 'application-default' in the Application field with 'any'.
    Returns (normalized_list, note_or_None)."""
    normalized = []
    substituted = False
    for value in raw_values:
        if value.strip().lower() in ("application-default", "app-default"):
            normalized.append("any")
            substituted = True
        else:
            normalized.append(value)
    note = None
    if substituted:
        note = "'application-default' isn't valid in the Application field (only in Service) - used 'any' instead"
    return normalized, note


def process_request(req: dict, existing_rules: list, taken_names: set, rulebase, fw, existing_services_by_name: dict, taken_service_names: set, known_tag_names: set, taken_tag_names: set, known_address_names: set, address_object_values: dict, virtual_routers_to_check: list, interface_zone_cache: dict, ip_zone_cache: dict, route_cache: dict, apply_changes: bool) -> dict:
    """Runs one Request-sheet row through the full pipeline. Returns a
    result dict for the Implementation sheet."""
    result = dict(req)
    result["Timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Normalize + self-consolidate once, so every step uses the same form.
    req_source_list, _ = merge_addresses([], normalize_address_list(split_values(req.get("Source"))))
    req_dest_list, _ = merge_addresses([], normalize_address_list(split_values(req.get("Destination"))))
    raw_service_values = build_raw_service_values(req.get("Service"), req.get("Destination Port"))
    req_application_list, application_note = normalize_application_list(split_values(req.get("Application")) or ["any"])

    # --- STEP 0: Resolve zones from routing ---
    print("  [1/5] Resolving zones from routing...")
    source_zone, source_zone_problems = resolve_zone_for_list(
        fw, req_source_list, virtual_routers_to_check, address_object_values, interface_zone_cache, ip_zone_cache, route_cache
    )
    dest_zone, dest_zone_problems = resolve_zone_for_list(
        fw, req_dest_list, virtual_routers_to_check, address_object_values, interface_zone_cache, ip_zone_cache, route_cache
    )
    all_zone_problems = [f"Source: {p}" for p in source_zone_problems] + [f"Destination: {p}" for p in dest_zone_problems]
    print(f"        Source zone: {source_zone or 'UNRESOLVED'}, Destination zone: {dest_zone or 'UNRESOLVED'}")

    if all_zone_problems and all(("no route found" in p) or ("DEFAULT route" in p) for p in all_zone_problems):
        result["Rule Name"] = ""
        result["Outcome"] = "Could not implement - no specific route found"
        result["Details"] = (
            "Could not be implemented because no SPECIFIC route exists on any checked virtual router "
            "(a match against the default route alone does not count): "
            + "; ".join(all_zone_problems)
        )
        print(f"        FAILED: {result['Details']}")
        return result

    # --- STEP 1: Validate objects ---
    print("  [2/5] Validating objects...")
    problems = all_zone_problems + validate_row(req, raw_service_values, known_address_names, existing_services_by_name)
    if problems:
        result["Rule Name"] = ""
        result["Outcome"] = "Validation error"
        result["Details"] = "; ".join(problems)
        print(f"        FAILED: {result['Details']}")
        return result

    resolved_service_names, service_notes = resolve_service_names(
        raw_service_values, existing_services_by_name, taken_service_names, fw, apply_changes
    )

    req_source_set = set(req_source_list)
    req_dest_set = set(req_dest_list)

    # --- STEP 2: Check existing rules ---
    print("  [3/5] Checking existing rules...")
    matched_rule, merge_field = find_matching_rule(
        existing_rules, req, resolved_service_names, req_application_list,
        source_zone, dest_zone, req_source_set, req_dest_set,
    )

    # --- STEP 3: Duplicate / merge / shadowing ---
    print("  [4/5] Checking for duplicates/merge/shadowing...")
    merged_values = None
    if matched_rule is not None:
        result["Rule Name"] = matched_rule.name
        changes = []
        if merge_field is not None:
            current_values = _rule_field(matched_rule, merge_field)
            if merge_field == "service":
                merged_values, changes = merge_service_names(current_values, resolved_service_names)
            else:
                new_values = req_source_list if merge_field == "source" else req_dest_list
                merged_values, changes = merge_addresses(current_values, new_values)

        if not changes:
            result["Outcome"] = "Duplicate - already exists"
            result["Details"] = f"Rule '{matched_rule.name}' already covers this exact request - nothing to do."
        else:
            label = {"source": "Source", "destination": "Destination", "service": "Service/Port"}[merge_field]
            result["Outcome"] = "Merged into existing rule"
            details = f"Updated {label} on rule '{matched_rule.name}': " + "; ".join(changes) + "."
            if service_notes:
                details += " Also created " + ", ".join(service_notes) + "."
            if application_note:
                details += " " + application_note + "."
            result["Details"] = details
    else:
        shadow_rule = check_shadowing(existing_rules, req, resolved_service_names, req_application_list, source_zone, dest_zone, req_source_set, req_dest_set)
        if shadow_rule is not None:
            result["Rule Name"] = ""
            result["Outcome"] = "Shadow conflict - not created"
            result["Details"] = (
                f"Existing rule '{shadow_rule.name}' (action={shadow_rule.action}) already matches this traffic "
                f"and sits earlier in the rulebase - a new rule below it would never be reached. "
                f"Not creating anything - review rule order manually."
            )
        else:
            result["Rule Name"] = None  # filled in below

    # --- STEP 4: Show proposed change / STEP 5: apply ---
    if result.get("Rule Name") is None:
        new_name = make_rule_name(req, dest_zone, taken_names)
        taken_names.add(new_name)

        resolved_tag_names, tag_notes = resolve_tags(
            split_values(req.get("Tag")), known_tag_names, taken_tag_names, fw, apply_changes
        )

        new_rule = SecurityRule(
            name=new_name,
            fromzone=[source_zone],
            tozone=[dest_zone],
            source=req_source_list,
            destination=req_dest_list,
            service=resolved_service_names,
            application=req_application_list,
            action=str(req.get("Action") or "").strip(),
            tag=resolved_tag_names,
        )

        result["Rule Name"] = new_name
        result["Outcome"] = "Created new rule"
        details = f"No existing rule matched, created '{new_name}' (Source Zone: {source_zone}, Destination Zone: {dest_zone})."
        if service_notes:
            details += " Also created " + ", ".join(service_notes) + "."
        if tag_notes:
            details += " Also created " + ", ".join(tag_notes) + "."
        if application_note:
            details += " " + application_note + "."
        result["Details"] = details

        # In-memory append is unconditional so a dry run simulates the batch.
        if apply_changes:
            rulebase.add(new_rule)
            new_rule.create()
        existing_rules.append(new_rule)
    elif result["Outcome"] == "Merged into existing rule":
        # In-memory update is unconditional (batch simulation); only the
        # device write is gated by apply_changes.
        setattr(matched_rule, merge_field, merged_values)
        if apply_changes:
            matched_rule.update(merge_field)

    print(f"  [5/5] Proposed change: {result['Outcome']} - {result['Details']}")
    return result


# --- Writing the Implementation sheet -------------------------------------

def write_results(excel_path: Path, sheet_name: str, results: list):
    wb = openpyxl.load_workbook(excel_path)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    columns = REQUEST_COLUMNS + ["Rule Name", "Outcome", "Details", "Timestamp"]
    ws.append(columns)
    for result in results:
        ws.append([result.get(col, "") for col in columns])

    wb.save(excel_path)


# --- Main ------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Validate and push Palo Alto security policies from an Excel sheet (firewall + credentials from AWX).")
    parser.add_argument("--excel", default=None, help="Path to the Excel file (default: excel_path from paloalto_config.json)")
    parser.add_argument("--request-sheet", default=None, help="Sheet name with the requested policies")
    parser.add_argument("--implementation-sheet", default=None, help="Sheet name to write results to")
    parser.add_argument("--apply", action="store_true", default=None, help="Actually create/update rules on the firewall")
    parser.add_argument("--commit", action="store_true", default=None, help="Commit the candidate config after applying")
    args = parser.parse_args()

    config = load_config()
    connection = load_awx_connection()

    excel_path = Path(args.excel if args.excel is not None else config["excel_path"])
    request_sheet = args.request_sheet if args.request_sheet is not None else config["request_sheet"]
    implementation_sheet = args.implementation_sheet if args.implementation_sheet is not None else config["implementation_sheet"]
    apply_changes = args.apply if args.apply is not None else config["apply"]
    do_commit = args.commit if args.commit is not None else config["commit"]

    if not excel_path.is_absolute():
        excel_path = get_base_dir() / excel_path

    print(f"Excel file: {excel_path}")
    if not excel_path.exists():
        print(f"ERROR: Excel file not found at {excel_path}")
        sys.exit(1)

    if do_commit and not apply_changes:
        print("NOTE: 'commit' is on but 'apply' is off - nothing to commit. Ignoring commit for this run.")
        do_commit = False

    print(f"Connecting to firewall {connection['hostname']} (from AWX inventory)...")
    fw, rulebase, existing_rules, existing_services_by_name, known_tag_names, zone_names, known_address_names, address_object_values, virtual_routers_to_check = connect_to_firewall(connection, config)
    taken_names = {rule.name for rule in existing_rules}
    taken_service_names = set()
    taken_tag_names = set()
    interface_zone_cache = {}
    ip_zone_cache = {}
    route_cache = {}
    print(f"Found {len(existing_rules)} existing security rule(s) on the firewall.")
    print(f"Found {len(existing_services_by_name)} existing service object(s), {len(known_tag_names)} tag(s), {len(zone_names)} zone(s), {len(known_address_names)} address object(s)/group(s) on the firewall.")
    print(f"Checking virtual router(s) for zone-from-routing lookups: {', '.join(virtual_routers_to_check)}")

    print(f"Reading requests from '{request_sheet}' sheet...")
    requests = read_requests(excel_path, request_sheet)
    print(f"Found {len(requests)} request row(s).")

    print(f"\nMODE: {'APPLY' if apply_changes else 'DRY RUN'}" + (" + COMMIT" if do_commit else ""))
    if not apply_changes:
        print("No changes will be made on the firewall.\n")
    else:
        print()

    results = []
    for i, req in enumerate(requests, start=1):
        print(f"Row {i}:")
        result = process_request(
            req, existing_rules, taken_names, rulebase, fw,
            existing_services_by_name, taken_service_names,
            known_tag_names, taken_tag_names,
            known_address_names, address_object_values, virtual_routers_to_check,
            interface_zone_cache, ip_zone_cache, route_cache,
            apply_changes=apply_changes,
        )
        results.append(result)

    print(f"\nWriting results to '{implementation_sheet}' sheet in {excel_path}...")
    write_results(excel_path, implementation_sheet, results)

    if apply_changes and do_commit:
        print("Committing candidate config to the firewall (this can take a minute)...")
        commit_result = fw.commit(sync=True)
        print(f"Commit result: {commit_result}")
    elif apply_changes:
        print("Changes were applied to the CANDIDATE config but NOT committed.")

    merged = sum(1 for r in results if r["Outcome"] == "Merged into existing rule")
    created = sum(1 for r in results if r["Outcome"] == "Created new rule")
    duplicates = sum(1 for r in results if r["Outcome"] == "Duplicate - already exists")
    shadowed = sum(1 for r in results if r["Outcome"] == "Shadow conflict - not created")
    no_route = sum(1 for r in results if r["Outcome"] == "Could not implement - no specific route found")
    invalid = sum(1 for r in results if r["Outcome"] == "Validation error")
    print(f"\nSummary: {merged} merged, {created} created, {duplicates} duplicate, {shadowed} shadow conflict(s), {no_route} no-route, {invalid} validation error(s).")


if __name__ == "__main__":
    main()
    # Pause only when a human is at a console (double-click / .exe).
    # AWX has no interactive stdin - input() there would crash the job.
    if sys.stdin is not None and sys.stdin.isatty():
        input("\nDone. Press Enter to close this window...")
